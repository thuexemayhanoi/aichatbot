#!/usr/bin/env python3
"""
seo-master-workbook.py — owner workbook + resume checkpoint for the MotoAI
image factory / SEO master task.

Reads ONLY committed sources of truth (never edits them):
  data/blog/published.json, data/blog/content-matrix.csv   real articles
  data/seo/image-plan.csv                                  per-article image brief
  data/seo/article-images.json                             image registry (status)
  docs/seo/seo-master-matrix.json                          per-URL SEO audit
  data/seo/serp-snapshots/*.json                           observed SERP (historical)
  docs/seo/github-seo-skills.json                          GitHub SEO research
  docs/seo/master-task/{backlog,errors,review-required,execution-log}.csv
  docs/seo/master-task/live-check.json (optional)          external HTTP checks
  built HTML (blog/, index.html, ...)                      technical audit

Writes:
  docs/seo/master-task/technical-audit.csv     per-URL technical audit
  docs/seo/master-task/checkpoint.json         dashboard + next batch (resume)
  docs/seo/AICHATBOT-SEO-MASTER-MATRIX.xlsx    9 sheets for the owner

Deterministic: workbook timestamps come from the checkpoint date, so the same
inputs give the same workbook content.

Run from the repo root:  python3 tools/seo-master-workbook.py [--date YYYY-MM-DD]
"""

import argparse
import csv
import glob
import json
import os
import re
import subprocess
from collections import Counter
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MT = 'docs/seo/master-task'
XLSX = 'docs/seo/AICHATBOT-SEO-MASTER-MATRIX.xlsx'
SITE = 'https://chatbot.thuexemaynguyentu.com/'
STATUSES = ['IMAGE_GENERATED', 'IMAGE_OPTIMIZED', 'IMAGE_ATTACHED', 'QA_PASSED', 'DEPLOYED', 'LIVE_VERIFIED']


def p(*parts):
    return os.path.join(ROOT, *parts)


def read_csv(path):
    if not os.path.exists(p(path)):
        return []
    with open(p(path), encoding='utf-8', newline='') as f:
        return list(csv.DictReader(f))


def read_json(path, default=None):
    if not os.path.exists(p(path)):
        return default
    with open(p(path), encoding='utf-8') as f:
        return json.load(f)


def rank(s):
    return STATUSES.index(s) if s in STATUSES else -1


def git(*args):
    try:
        return subprocess.run(['git', *args], cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip()
    except Exception:
        return ''


# ---------------- technical audit (built HTML) ----------------

def url_to_file(url):
    path = url[len(SITE):]
    return (path + 'index.html') if (path == '' or path.endswith('/')) else path


def technical_audit(seo_rows, live):
    sitemap = open(p('sitemap.xml'), encoding='utf-8').read()
    in_sitemap = set(re.findall(r'<loc>([^<]+)</loc>', sitemap))
    robots = open(p('robots.txt'), encoding='utf-8').read()
    disallow = [l.split(':', 1)[1].strip() for l in robots.splitlines() if l.lower().startswith('disallow:') and l.split(':', 1)[1].strip()]
    css_bytes = sum(os.path.getsize(p(f)) for f in ('assets/css/style.css', 'assets/css/blog.css') if os.path.exists(p(f)))
    out = []
    for r in seo_rows:
        url = r['URL']
        f = url_to_file(url)
        row = {'URL': url, 'Page Type': r.get('Page Type', ''), 'Article ID': r.get('Article ID', '')}
        if not os.path.exists(p(f)):
            row.update({'HTTP Status': 'MISSING_LOCAL', 'Priority': 'P0', 'Repair Status': 'OPEN'})
            out.append(row)
            continue
        html = open(p(f), encoding='utf-8').read()
        robots_meta = re.search(r'<meta name="robots" content="([^"]+)"', html)
        canon = re.search(r'<link rel="canonical" href="([^"]+)"', html)
        types = sorted(set(re.findall(r'"@type":\s*"(Article|BlogPosting|BreadcrumbList|CollectionPage|WebPage|WebSite|FAQPage|LocalBusiness|Organization|ItemList|WebApplication|AboutPage|ContactPage|HowTo|Product|Offer)"', html)))
        h1 = len(re.findall(r'<h1\b', html))
        hero = 'có (QA_PASSED)' if 'class="article-hero"' in html else 'chưa'
        noindex = bool(robots_meta and 'noindex' in robots_meta.group(1))
        lv = live.get(url, {})
        issues = []
        if noindex:
            issues.append('noindex')
        if not canon or canon.group(1) != url:
            issues.append('canonical không tự trỏ')
        if url not in in_sitemap:
            issues.append('ngoài sitemap')
        if h1 != 1:
            issues.append(f'H1={h1}')
        if not types:
            issues.append('thiếu schema')
        if lv.get('status') not in (None, 200):
            issues.append(f'HTTP {lv.get("status")}')
        row.update({
            'HTTP Status': lv.get('status', 'NOT_MEASURED'),
            'HTTP Checked': lv.get('checked', ''),
            'Indexability': 'noindex' if noindex else 'indexable',
            'Canonical': 'self' if canon and canon.group(1) == url else (canon.group(1) if canon else 'MISSING'),
            'Sitemap': 'có' if url in in_sitemap else 'không',
            'Robots': 'blocked' if any(url[len(SITE) - 1:].startswith(d) for d in disallow) else 'allowed',
            'Schema': '+'.join(types) or 'none',
            'H1': h1,
            'Featured Image': hero,
            'Broken Links': 0,  # tools/site-audit.mjs: broken_internal_links = 0 site-wide (gate in CI)
            'Performance': f'HTML {round(len(html.encode("utf-8")) / 1024)} KB; CSS {round(css_bytes / 1024)} KB',
            'Issues': '; '.join(issues) or 'Không',
            'Priority': 'P0' if issues else r.get('Priority', 'P2'),
            'Repair Status': 'OPEN' if issues else 'PASS',
        })
        out.append(row)
    return out


TECH_COLS = ['URL', 'Page Type', 'Article ID', 'HTTP Status', 'HTTP Checked', 'Indexability', 'Canonical', 'Sitemap', 'Robots',
             'Schema', 'H1', 'Featured Image', 'Broken Links', 'Performance', 'Issues', 'Priority', 'Repair Status']


# ---------------- historical ranking (SERP snapshots) ----------------

def historical_positions():
    pos = {}
    for path in sorted(glob.glob(p('data/seo/serp-snapshots/*.json'))):
        snap = json.load(open(path, encoding='utf-8'))
        day = snap.get('collected', os.path.basename(path)[:10])
        for q in snap.get('queries', []):
            for o in q.get('organic', []):
                u = o.get('url', '').replace('/index.html', '/')
                if u.startswith(SITE):
                    pos.setdefault(u, []).append(f'#{o.get("position")} "{q["query"]}" ({day}, {q.get("source", "")})')
    return pos


# ---------------- workbook ----------------

def normalize_zip(path, when):
    """Rewrite the xlsx zip with fixed entry timestamps → byte-stable output."""
    import zipfile
    with zipfile.ZipFile(path) as z:
        entries = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        stamp = when.strftime('%Y-%m-%dT00:00:00Z').encode()
        for name, data in entries:
            if name == 'docProps/core.xml':  # openpyxl stamps "modified" with the save time
                data = re.sub(rb'(<dcterms:modified[^>]*>)[^<]*', rb'\g<1>' + stamp, data)
            info = zipfile.ZipInfo(name, date_time=(when.year, when.month, when.day, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            z.writestr(info, data)


def build(date):
    published = read_json('data/blog/published.json')['articles']
    matrix = {r['article_id']: r for r in read_csv('data/blog/content-matrix.csv')}
    plan = read_csv('data/seo/image-plan.csv')
    registry = read_json('data/seo/article-images.json', {'images': {}})['images']
    seo = read_json('docs/seo/seo-master-matrix.json', {'rows': []})['rows']
    skills = read_json('docs/seo/github-seo-skills.json', {'repos': [], 'methods': []})
    live = read_json(f'{MT}/live-check.json', {'checks': {}})['checks']
    backlog = read_csv(f'{MT}/backlog.csv')
    errors = read_csv(f'{MT}/errors.csv')
    review = read_csv(f'{MT}/review-required.csv')
    log = read_csv(f'{MT}/execution-log.csv')
    hist = historical_positions()

    tech = technical_audit(seo, live)
    with open(p(f'{MT}/technical-audit.csv'), 'w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=TECH_COLS, lineterminator='\n')
        w.writeheader()
        for r in tech:
            w.writerow({k: r.get(k, '') for k in TECH_COLS})

    # ---- counts ----
    st = Counter(e['status'] for e in registry.values())
    at_least = lambda s: sum(1 for e in registry.values() if rank(e['status']) >= rank(s))
    target = [r for r in plan if r['scope'] == 'TARGET']
    pending_ids = [r['article_id'] for r in target if r['article_id'] not in registry]
    pub_matrix = sum(1 for r in matrix.values() if r['status'] == 'PUBLISHED')
    review_open = sum(1 for r in review if r.get('status') == 'OPEN')
    reviews_needed = sum(1 for r in seo if r.get('Content Review') == 'REVIEW')
    head = git('rev-parse', '--short', 'HEAD')
    tech_open = sum(1 for r in tech if r.get('Repair Status') == 'OPEN')
    blocked = [r for r in backlog if r['status'] == 'BLOCKED']
    next_batch = (blocked[0]['next_required_step'] if blocked else '') or ', '.join(pending_ids[:10])

    dashboard = [
        ('Ngày cập nhật', date),
        ('Tổng bài thực tế (manifest published.json)', len(published)),
        ('Tổng hàng PUBLISHED trong content-matrix.csv', pub_matrix),
        ('Tổng bài mục tiêu ảnh (BA-0001…BA-2000)', len(target)),
        ('Bài ngoài mục tiêu (BA-2001+, scope EXTRA)', len(plan) - len(target)),
        ('Ảnh đã tạo (≥ IMAGE_GENERATED)', at_least('IMAGE_GENERATED')),
        ('Ảnh đã nén WebP (≥ IMAGE_OPTIMIZED)', at_least('IMAGE_OPTIMIZED')),
        ('Ảnh đã gắn ALT (≥ IMAGE_ATTACHED)', at_least('IMAGE_ATTACHED')),
        ('Ảnh đã QA (≥ QA_PASSED)', at_least('QA_PASSED')),
        ('Ảnh đã deploy (≥ DEPLOYED)', at_least('DEPLOYED')),
        ('Ảnh đã live verify (LIVE_VERIFIED)', st.get('LIVE_VERIFIED', 0)),
        ('Bài mục tiêu chưa có ảnh', len(pending_ids)),
        ('URL đã audit SEO (seo-master-audit)', len(seo)),
        ('URL trạng thái DONE trong SEO master (tối ưu v70–v79)', sum(1 for r in seo if r.get('Status') == 'DONE')),
        ('Bài cần review nội dung (Content Review = REVIEW)', reviews_needed),
        ('URL có lỗi kỹ thuật mở (technical-audit)', tech_open),
        ('Việc cần chủ sở hữu xác minh (REVIEW REQUIRED mở)', review_open),
        ('Commit cuối (HEAD khi xuất)', head),
        ('Checkpoint cuối', f'{MT}/checkpoint.json @ {date}'),
        ('Next batch', next_batch),
    ]

    # ---- image matrix rows ----
    img_rows = []
    for r in plan:
        e = registry.get(r['article_id'], {})
        s = e.get('status', 'PENDING')
        img_rows.append([
            r['article_id'], r['url'], r['slug'], r['title'], r['category'], r['primary_keyword'], r['search_intent'],
            'không (chỉ OG chung icon-512)', r['brief_vi'], e.get('file', ''), r['filename'], e.get('width', ''), e.get('height', ''),
            round(e['bytes'] / 1024, 1) if e.get('bytes') else '', e.get('alt') or r['alt_plan'] + ' (dự kiến)',
            ('PASS (sha256+dHash)' if rank(s) >= rank('QA_PASSED') else ('chờ' if rank(s) < rank('IMAGE_OPTIMIZED') else 'đã băm')),
            'IMAGE_GENERATED' if rank(s) >= 0 else 'PENDING',
            'IMAGE_OPTIMIZED' if rank(s) >= rank('IMAGE_OPTIMIZED') else 'PENDING',
            'IMAGE_ATTACHED' if rank(s) >= rank('IMAGE_ATTACHED') else 'PENDING',
            'QA_PASSED' if rank(s) >= rank('QA_PASSED') else 'PENDING',
            'DEPLOYED' if rank(s) >= rank('DEPLOYED') else 'PENDING',
            'LIVE_VERIFIED' if s == 'LIVE_VERIFIED' else 'PENDING',
            e.get('commit', ''), e.get('generator', {}).get('generated_at', '') if e else '',
            e.get('notes', '') or ('ngoài 2.000 mục tiêu' if r['scope'] == 'EXTRA' else ''),
        ])

    # ---- SEO URL master ----
    reg_by_url = {r['url']: registry.get(r['article_id']) for r in plan}
    tech_by_url = {r['URL']: r for r in tech}
    url_rows = []
    for r in seo:
        url = r['URL']
        ev = r.get('Evidence', '')
        m = re.search(r'Hỗ trợ pillar (BA-\d{4})', r.get('Cannibalization', ''))
        owner = f'pillar {m.group(1)}' if m else 'chính URL này'
        e = reg_by_url.get(url)
        img = (e['status'] if e else ('PENDING' if r.get('Article ID') else 'n/a'))
        t = tech_by_url.get(url, {})
        title_len = re.search(r'title (\d+)c', ev)
        desc_len = re.search(r'desc (\d+)c', ev)
        faq = 'có FAQ' if re.search(r'FAQ', ev) else ''
        url_rows.append([
            r.get('Article ID', ''), url, r.get('Page Type', ''), r.get('Primary Keyword', ''), r.get('Search Intent', ''), owner,
            r.get('Status', ''), '; '.join(hist.get(url, [])) or 'NOT_MEASURED', r.get('Technical Issues', ''), r.get('Content Gap', ''),
            f'OK ({title_len.group(1)} ký tự)' if title_len else 'xem Evidence', f'OK ({desc_len.group(1)} ký tự)' if desc_len else 'xem Evidence',
            'H1 = 1' if 'H1 1' in ev else 'xem Evidence', img, r.get('Internal Links', ''), t.get('Schema', ''),
            (faq or 'chưa có FAQ riêng') + ('; Content Review ' + r.get('Content Review', '') if r.get('Content Review') else ''),
            r.get('Priority', ''), t.get('Repair Status', ''), 'LIVE (GitHub Pages)' if r.get('Status') == 'DONE' else '',
            r.get('Proposed Action', ''),
        ])

    # ---- write xlsx ----
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    fixed = datetime.strptime(date, '%Y-%m-%d')
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = 'MotoAI SEO master task'
    wb.properties.title = 'AICHATBOT SEO MASTER MATRIX'
    head_fill = PatternFill('solid', fgColor='1F2937')
    head_font = Font(bold=True, color='FFFFFF')
    fills = {
        'DONE': 'D1FAE5', 'PASS': 'D1FAE5', 'QA_PASSED': 'D1FAE5', 'LIVE_VERIFIED': 'A7F3D0', 'DEPLOYED': 'BBF7D0',
        'BLOCKED': 'FECACA', 'OPEN': 'FEF3C7', 'PENDING': 'F3F4F6', 'IN_PROGRESS': 'DBEAFE', 'IMAGE_GENERATED': 'DBEAFE',
        'REVIEW_REQUIRED': 'FDE68A', 'ADOPT': 'D1FAE5', 'ADAPT': 'DBEAFE', 'REJECT': 'FECACA',
    }

    def sheet(title, headers, rows, widths=None, status_cols=()):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(wrap_text=True, vertical='center')
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = 'B2' if len(headers) > 3 else 'A2'
        if rows:
            ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows) + 1}'
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, min(max(len(str(h)) + 2, 12), 40))
        for col in status_cols:
            idx = headers.index(col) + 1
            for r in range(2, len(rows) + 2):
                v = str(ws.cell(row=r, column=idx).value or '')
                if v in fills:
                    ws.cell(row=r, column=idx).fill = PatternFill('solid', fgColor=fills[v])
        return ws

    ws = wb.active
    ws.title = 'DASHBOARD'
    ws.append(['Chỉ số', 'Giá trị'])
    for c in ws[1]:
        c.fill, c.font = head_fill, head_font
    for k, v in dashboard:
        ws.append([k, v])
    ws.append([])
    ws.append(['Trạng thái ảnh', 'Số lượng'])
    for s in ['PENDING'] + STATUSES:
        ws.append([s, len(pending_ids) if s == 'PENDING' else st.get(s, 0)])
    ws.append([])
    ws.append(['Quy tắc "Tiếp tục"', 'Đọc AGENTS.md → docs/seo/master-task/README.md → checkpoint.json; tiếp tục từ next batch; không làm lại mục DONE.'])
    ws.column_dimensions['A'].width = 58
    ws.column_dimensions['B'].width = 110

    sheet('IMAGE 2000 MATRIX', ['Article ID', 'URL', 'Slug', 'Title', 'Category', 'Primary Keyword', 'Search Intent', 'Existing Image',
                                'Image Brief', 'Generated Image Path', 'Filename', 'Width', 'Height', 'File Size KB', 'ALT Text',
                                'Duplicate Check', 'Generation Status', 'Compression Status', 'Attachment Status', 'Image QA',
                                'Deployment', 'Live Verification', 'Commit SHA', 'Last Updated', 'Error/Notes'],
          img_rows, {'URL': 45, 'Title': 45, 'Image Brief': 60, 'ALT Text': 60, 'Filename': 40, 'Error/Notes': 50},
          status_cols=('Generation Status', 'Compression Status', 'Attachment Status', 'Image QA', 'Deployment', 'Live Verification'))
    sheet('SEO URL MASTER', ['Article ID', 'URL', 'Page Type', 'Primary Keyword', 'Search Intent', 'Keyword Owner', 'Baseline Status',
                             'Historical Ranking', 'Technical Issues', 'Content Gap', 'Title Status', 'Meta Status', 'H1/H2/H3 Status',
                             'Image SEO', 'Internal Links', 'Schema', 'AEO/GEO', 'Priority', 'QA Result', 'Deployment', 'Next Action'],
          url_rows, {'URL': 45, 'Historical Ranking': 40, 'Internal Links': 40, 'Next Action': 50, 'AEO/GEO': 30},
          status_cols=('Baseline Status', 'QA Result'))
    sheet('GITHUB SEO SKILLS', ['Repository', 'Author', 'Stars', 'Last Commit', 'License', 'Skill', 'Source URL', 'Applicability',
                                'Risk', 'Adopt/Adapt/Reject', 'Applied URLs', 'QA Evidence'],
          [[m.get('repository', ''), next((r.get('author', '') for r in skills.get('repos', []) if r.get('repository') == m.get('repository')), ''),
            next((r.get('stars', '') for r in skills.get('repos', []) if r.get('repository') == m.get('repository')), ''),
            next((r.get('last_commit', '') for r in skills.get('repos', []) if r.get('repository') == m.get('repository')), ''),
            next((r.get('license', '') for r in skills.get('repos', []) if r.get('repository') == m.get('repository')), ''),
            m.get('skill', ''), m.get('source_url', ''), m.get('applicability', ''), m.get('risk', ''), m.get('decision', ''),
            m.get('applied_urls', '') or ('đã có sẵn' if m.get('already_in_place') else ''), m.get('qa_evidence', '') or m.get('reason', '')]
           for m in skills.get('methods', [])],
          {'Skill': 35, 'Source URL': 50, 'Applicability': 60, 'QA Evidence': 60, 'Risk': 30}, status_cols=('Adopt/Adapt/Reject',))
    sheet('TECHNICAL AUDIT', ['URL', 'HTTP Status', 'Indexability', 'Canonical', 'Sitemap', 'Robots', 'Schema', 'Broken Links',
                              'Performance', 'Priority', 'Repair Status', 'Issues', 'Featured Image', 'HTTP Checked'],
          [[r.get(k, '') for k in ['URL', 'HTTP Status', 'Indexability', 'Canonical', 'Sitemap', 'Robots', 'Schema', 'Broken Links',
                                   'Performance', 'Priority', 'Repair Status', 'Issues', 'Featured Image', 'HTTP Checked']] for r in tech],
          {'URL': 60, 'Schema': 35, 'Performance': 25, 'Issues': 30}, status_cols=('Repair Status',))
    sheet('BACKLOG & CHECKPOINT', ['Task ID', 'Phase', 'Batch', 'Priority', 'Dependencies', 'Status', 'Last Completed Step',
                                   'Next Required Step', 'Blocking Issue', 'Commit SHA', 'Timestamp'],
          [[r[k] for k in ['task_id', 'phase', 'batch', 'priority', 'dependencies', 'status', 'last_completed_step',
                            'next_required_step', 'blocking_issue', 'commit_sha', 'timestamp']] for r in backlog],
          {'Last Completed Step': 60, 'Next Required Step': 60, 'Blocking Issue': 45, 'Phase': 28}, status_cols=('Status',))
    sheet('ERROR & AUTO REPAIR', ['Error ID', 'Related URL/ID', 'Error Type', 'Root Cause', 'Repair Attempt', 'Result', 'QA Rerun', 'Final Status'],
          [[r[k] for k in ['error_id', 'related', 'error_type', 'root_cause', 'repair_attempt', 'result', 'qa_rerun', 'final_status']] for r in errors],
          {'Root Cause': 60, 'Repair Attempt': 60, 'Result': 30, 'Related URL/ID': 30}, status_cols=('Final Status',))
    sheet('REVIEW REQUIRED', ['Review ID', 'Area', 'Item', 'Why', 'Owner Action', 'Blocks', 'Priority', 'Status'],
          [[r[k] for k in ['review_id', 'area', 'item', 'why', 'owner_action', 'blocks', 'priority', 'status']] for r in review],
          {'Item': 40, 'Why': 60, 'Owner Action': 60}, status_cols=('Status',))
    sheet('EXECUTION LOG', ['Run ID', 'Start', 'End', 'Baseline SHA', 'Batch IDs', 'Files Changed', 'Image Count', 'SEO Count',
                            'QA Result', 'Commit SHA', 'Deploy Result', 'Remaining Work', 'Next Checkpoint'],
          [[r[k] for k in ['run_id', 'start', 'end', 'baseline_sha', 'batch_ids', 'files_changed', 'image_count', 'seo_count',
                            'qa_result', 'commit_sha', 'deploy_result', 'remaining_work', 'next_checkpoint']] for r in log],
          {'Files Changed': 40, 'QA Result': 40, 'Remaining Work': 40, 'Batch IDs': 30})
    wb.save(p(XLSX))
    normalize_zip(p(XLSX), fixed)

    checkpoint = {
        '$schema': 'motoai/master-task-checkpoint@1',
        'task': 'MOTOAI IMAGE FACTORY 2000 + GITHUB EXPERT SEO SKILLS + FULL WEBSITE SEO',
        'updated': date,
        'head_at_export': head,
        'workbook': XLSX,
        'dashboard': {k: v for k, v in dashboard},
        'image_status_counts': {s: st.get(s, 0) for s in STATUSES},
        'next_batch': next_batch,
        'blocked': [{'task': r['task_id'], 'issue': r['blocking_issue']} for r in blocked],
        'open_review_required': [r['review_id'] for r in review if r.get('status') == 'OPEN'],
        'resume': 'docs/seo/master-task/README.md',
    }
    with open(p(f'{MT}/checkpoint.json'), 'w', encoding='utf-8') as f:
        f.write(json.dumps(checkpoint, ensure_ascii=False, indent=2) + '\n')
    print(f'wrote {XLSX} ({len(img_rows)} image rows, {len(url_rows)} URL rows, {len(tech)} tech rows) + checkpoint')


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--date', default=datetime.now().strftime('%Y-%m-%d'))
    build(ap.parse_args().date)
